"""Document processing utilities for RAG system.

This module handles text extraction, chunking, and preprocessing for various file formats.
"""

import mimetypes
from pathlib import Path
from typing import List, Tuple, Optional
from loguru import logger

# Text extraction libraries
try:
    import pypdf

    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    logger.warning("pypdf not installed. PDF processing will be disabled.")

try:
    import docx

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not installed. DOCX processing will be disabled.")

try:
    from bs4 import BeautifulSoup

    HTML_AVAILABLE = True
except ImportError:
    HTML_AVAILABLE = False
    logger.warning("beautifulsoup4 not installed. HTML processing will be disabled.")


class DocumentProcessor:
    """Handles document text extraction and processing."""

    SUPPORTED_FORMATS = {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".csv": "text/csv",
        ".json": "application/json",
        ".xml": "application/xml",
        ".html": "text/html",
        ".htm": "text/html",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }

    @staticmethod
    def detect_file_type(file_path: str) -> str:
        """Detect file type from extension and MIME type.

        Args:
            file_path: Path to the file

        Returns:
            File extension (e.g., '.pdf')
        """
        ext = Path(file_path).suffix.lower()
        if ext in DocumentProcessor.SUPPORTED_FORMATS:
            return ext

        # Fallback to MIME type detection
        mime_type, _ = mimetypes.guess_type(file_path)
        for known, supported_mime in DocumentProcessor.SUPPORTED_FORMATS.items():
            if mime_type == supported_mime:
                return known

        # Named from the file, not the loop: reusing `ext` as the loop variable
        # reported every unsupported file (.exe, no extension) as ".docx".
        supported = ", ".join(sorted(DocumentProcessor.SUPPORTED_FORMATS))
        raise ValueError(
            f"Unsupported file type: {ext or 'no extension'} ({Path(file_path).name}). "
            f"Supported: {supported}"
        )

    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        """Extract text from plain text files."""
        encodings = ["utf-8", "latin-1", "cp1252"]

        for encoding in encodings:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    return f.read()
            except (UnicodeDecodeError, LookupError):
                continue

        raise ValueError(f"Could not decode text file with supported encodings: {file_path}")

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Extract text from PDF files.

        Handles various PDF types including those with unusual structures.
        Falls back to pdfplumber if pypdf fails.
        """
        if not PDF_AVAILABLE:
            raise ImportError(
                "pypdf is required for PDF processing. Install with: pip install pypdf"
            )

        text = []
        extraction_errors = []

        try:
            with open(file_path, "rb") as f:
                pdf_reader = pypdf.PdfReader(f)
                num_pages = len(pdf_reader.pages)
                logger.debug(f"PDF has {num_pages} pages")

                for page_num in range(num_pages):
                    try:
                        page = pdf_reader.pages[page_num]
                        page_text = page.extract_text()
                        if page_text:
                            text.append(f"[Page {page_num + 1}]\n{page_text}")
                    except Exception as page_error:
                        # Log but continue with other pages
                        extraction_errors.append(f"Page {page_num + 1}: {str(page_error)}")
                        logger.warning(
                            f"Failed to extract text from page {page_num + 1}: {page_error}"
                        )
                        continue

        except Exception as e:
            logger.error(f"Error reading PDF {file_path}: {e}")
            # Try fallback with pdfplumber if available
            try:
                import pdfplumber

                logger.info(f"Trying pdfplumber fallback for {file_path}")
                with pdfplumber.open(file_path) as pdf:
                    for page_num, page in enumerate(pdf.pages):
                        page_text = page.extract_text()
                        if page_text:
                            text.append(f"[Page {page_num + 1}]\n{page_text}")
            except ImportError:
                logger.warning("pdfplumber not available for fallback")
                raise ValueError(
                    f"Failed to extract text from PDF: {str(e)}. Install pdfplumber for better PDF support: pip install pdfplumber"
                )
            except Exception as fallback_error:
                logger.error(f"Fallback PDF extraction also failed: {fallback_error}")
                raise ValueError(f"Failed to extract text from PDF: {str(e)}")

        if not text:
            if extraction_errors:
                raise ValueError(
                    f"Could not extract any text from PDF. Errors: {'; '.join(extraction_errors[:3])}"
                )
            raise ValueError("PDF appears to be empty or contains only images/scanned content")

        if extraction_errors:
            logger.warning(f"PDF extraction completed with {len(extraction_errors)} page errors")

        return "\n\n".join(text)

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from DOCX files."""
        if not DOCX_AVAILABLE:
            raise ImportError(
                "python-docx is required for DOCX processing. Install with: pip install python-docx"
            )

        try:
            doc = docx.Document(file_path)
            paragraphs = []
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = getattr(paragraph.style, "name", "") if paragraph.style else ""
                if style_name.startswith("Heading"):
                    level_str = style_name.replace("Heading", "").strip()
                    try:
                        level = max(1, min(6, int(level_str)))
                    except ValueError:
                        level = 1
                    text = f"{'#' * level} {text}"
                paragraphs.append(text)
            return "\n\n".join(paragraphs)
        except Exception as e:
            logger.error(f"Error extracting text from DOCX {file_path}: {e}")
            raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

    @staticmethod
    def extract_text_from_html(file_path: str) -> str:
        """Extract text from HTML files."""
        if not HTML_AVAILABLE:
            # Fallback to simple tag stripping
            with open(file_path, "r", encoding="utf-8") as f:
                html_content = f.read()
            # Simple tag removal (not ideal but works without dependencies)
            import re

            for level in range(6, 0, -1):
                pattern = re.compile(rf"<h{level}[^>]*>(.*?)</h{level}>", re.IGNORECASE | re.DOTALL)
                html_content = pattern.sub(lambda m: f"{'#' * level} {m.group(1)}", html_content)
            text = re.sub(r"<[^>]+>", " ", html_content)
            text = re.sub(r"\s+", " ", text)
            return text.strip()

        try:
            with open(file_path, "r", encoding="utf-8") as f:
                soup = BeautifulSoup(f.read(), "html.parser")

            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()

            for level in range(6, 0, -1):
                for heading in soup.find_all(f"h{level}"):
                    heading.replace_with(f"{'#' * level} {heading.get_text(strip=True)}")

            text = soup.get_text(separator="\n")
            # Clean up whitespace
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            return "\n".join(lines)

        except Exception as e:
            logger.error(f"Error extracting text from HTML {file_path}: {e}")
            raise ValueError(f"Failed to extract text from HTML: {str(e)}")

    @staticmethod
    def extract_text(file_path: str, file_type: Optional[str] = None) -> str:
        """Extract text from a document.

        Args:
            file_path: Path to the file
            file_type: Optional file type (extension). If not provided, will be detected.

        Returns:
            Extracted text content

        Raises:
            ValueError: If file type is unsupported or extraction fails
        """
        if file_type is None:
            file_type = DocumentProcessor.detect_file_type(file_path)

        logger.debug(f"Extracting text from {file_type} file: {file_path}")

        # Route to appropriate extractor
        if file_type in [".txt", ".md", ".csv", ".json", ".xml"]:
            return DocumentProcessor.extract_text_from_txt(file_path)
        elif file_type == ".pdf":
            return DocumentProcessor.extract_text_from_pdf(file_path)
        elif file_type == ".docx":
            return DocumentProcessor.extract_text_from_docx(file_path)
        elif file_type in [".html", ".htm"]:
            return DocumentProcessor.extract_text_from_html(file_path)
        else:
            raise ValueError(f"No extractor available for file type: {file_type}")


class TextChunker:
    """Splits text into overlapping chunks for embedding."""

    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200):
        """Initialize the chunker.

        Args:
            chunk_size: Maximum characters per chunk
            chunk_overlap: Number of characters to overlap between chunks
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def split_text(self, text: str) -> List[str]:
        """Split text into overlapping chunks.

        Args:
            text: The text to split

        Returns:
            List of text chunks
        """
        if not text or not text.strip():
            return []

        chunks = []
        start = 0
        text_length = len(text)

        while start < text_length:
            # Find the end of this chunk
            end = start + self.chunk_size

            # If this is not the last chunk, try to break at a sentence boundary
            if end < text_length:
                # Look for sentence endings in the last 20% of the chunk
                search_start = end - int(self.chunk_size * 0.2)
                chunk_text = text[search_start:end]

                # Try to find sentence boundaries
                for delimiter in [". ", ".\n", "! ", "!\n", "? ", "?\n"]:
                    last_sentence_end = chunk_text.rfind(delimiter)
                    if last_sentence_end != -1:
                        end = search_start + last_sentence_end + len(delimiter)
                        break

            # Extract the chunk
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)

            # Move to next chunk with overlap
            start = end - self.chunk_overlap

            # Prevent infinite loop if overlap is too large
            if start <= end - self.chunk_size:
                start = end

        logger.info(
            f"Split text into {len(chunks)} chunks (size={self.chunk_size}, overlap={self.chunk_overlap})"
        )
        return chunks

    def create_chunks_with_metadata(
        self, text: str, document_id: str, session_id: str, filename: str
    ) -> List[Tuple[str, dict]]:
        """Create chunks with associated metadata.

        Args:
            text: The text to split
            document_id: ID of the parent document
            session_id: Session this document belongs to
            filename: Original filename

        Returns:
            List of (chunk_text, metadata) tuples
        """
        chunks = self.split_text(text)
        result = []

        for idx, chunk_text in enumerate(chunks):
            metadata = {
                "document_id": document_id,
                "session_id": session_id,
                "filename": filename,
                "chunk_index": idx,
                "total_chunks": len(chunks),
            }
            result.append((chunk_text, metadata))

        return result


# ---------------------------------------------------------------------------
# ABC-compliant adapter
# ---------------------------------------------------------------------------

from lasagna.core.base_components import (  # noqa: E402
    BaseDocumentProcessor,
    ProcessedDocument,
)


class DocumentProcessorAdapter(BaseDocumentProcessor):
    """Adapts the static ``DocumentProcessor`` to the ``BaseDocumentProcessor`` ABC."""

    def extract(self, file_path: str, **kwargs) -> ProcessedDocument:
        # Only forward the one kwarg that extract_text accepts
        file_type = kwargs.get("file_type")
        text = DocumentProcessor.extract_text(file_path, file_type=file_type)
        return ProcessedDocument(
            text=text,
            metadata={"source": file_path},
        )

    @property
    def supported_extensions(self) -> List[str]:
        return list(DocumentProcessor.SUPPORTED_FORMATS.keys())
